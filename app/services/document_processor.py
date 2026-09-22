"""Document processing service.

Handles file validation, type detection, content extraction, and
preparation of multimodal inputs for the Gemini API.

Processing strategies:
- PDF  → sent as Gemini `document` type (native vision understanding)
- Image → sent as Gemini `image` type (multimodal vision)
- DOCX → text extracted via python-docx, sent as text
- RTF  → text extracted via striprtf, sent as text
- TXT/MD → read directly as text
"""

from __future__ import annotations

import base64
import io
import logging
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from app.config import (
    INLINE_SIZE_LIMIT_BYTES,
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB,
    SUPPORTED_EXTENSIONS,
    SUPPORTED_FILE_TYPES,
    SUPPORTED_FORMATS_DISPLAY,
)

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------

@dataclass
class ProcessedDocument:
    """Result of document processing — ready for Gemini API input."""
    file_name: str
    file_size: int
    content_type: str  # "document", "image", "text"
    mime_type: str

    # For document/image: base64-encoded bytes
    data_base64: Optional[str] = None

    # For text: extracted plain text
    text_content: Optional[str] = None

    # For Files API (large files): URI from uploaded file
    file_uri: Optional[str] = None

    # Processing metadata
    warnings: list[str] = field(default_factory=list)
    is_valid: bool = True
    error: Optional[str] = None


@dataclass
class ValidationResult:
    """Result of file validation."""
    is_valid: bool
    error: Optional[str] = None
    extension: str = ""
    mime_type: str = ""
    processing_strategy: str = ""


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def validate_file(file_name: str, file_size: int, file_bytes: bytes) -> ValidationResult:
    """Validate an uploaded file.

    Checks:
    1. File extension is supported
    2. File size is within limits
    3. File is not empty
    4. Basic content validation
    """
    # Extract extension
    ext = Path(file_name).suffix.lower()

    # Check extension
    if ext not in SUPPORTED_FILE_TYPES:
        return ValidationResult(
            is_valid=False,
            error=(
                f"Unsupported file format: '{ext}'. "
                f"Supported formats: {SUPPORTED_FORMATS_DISPLAY}"
            ),
        )

    # Check size
    if file_size > MAX_FILE_SIZE_BYTES:
        return ValidationResult(
            is_valid=False,
            error=f"File too large ({file_size / (1024*1024):.1f} MB). Maximum: {MAX_FILE_SIZE_MB} MB",
        )

    # Check empty
    if file_size == 0 or not file_bytes:
        return ValidationResult(
            is_valid=False,
            error="File is empty.",
        )

    # Basic content validation — check magic bytes for common types
    content_issue = _validate_content(ext, file_bytes)
    if content_issue:
        return ValidationResult(is_valid=False, error=content_issue)

    _, mime_type, strategy = SUPPORTED_FILE_TYPES[ext]

    return ValidationResult(
        is_valid=True,
        extension=ext,
        mime_type=mime_type,
        processing_strategy=strategy,
    )


def _validate_content(ext: str, data: bytes) -> Optional[str]:
    """Basic content validation using magic bytes.

    Returns an error message if the content appears corrupt, or None if valid.
    """
    if ext == ".pdf":
        if not data[:5].startswith(b"%PDF"):
            return "File does not appear to be a valid PDF (missing PDF header)."

    elif ext == ".docx":
        # DOCX is a ZIP file
        if not data[:4] == b"PK\x03\x04":
            return "File does not appear to be a valid DOCX file."

    elif ext in (".png",):
        if not data[:8] == b"\x89PNG\r\n\x1a\n":
            return "File does not appear to be a valid PNG image."

    elif ext in (".jpg", ".jpeg"):
        if not data[:2] == b"\xff\xd8":
            return "File does not appear to be a valid JPEG image."

    elif ext == ".webp":
        if not (data[:4] == b"RIFF" and data[8:12] == b"WEBP"):
            return "File does not appear to be a valid WebP image."

    return None


# ---------------------------------------------------------------------------
# Processing
# ---------------------------------------------------------------------------

def process_document(
    file_name: str,
    file_bytes: bytes,
    gemini_client=None,
) -> ProcessedDocument:
    """Process an uploaded file into a format ready for Gemini API.

    Args:
        file_name: Original uploaded file name.
        file_bytes: Raw file bytes.
        gemini_client: Optional Gemini client for Files API upload (large files).

    Returns:
        ProcessedDocument with either data_base64, text_content, or file_uri set.
    """
    file_size = len(file_bytes)

    # Validate
    validation = validate_file(file_name, file_size, file_bytes)
    if not validation.is_valid:
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="unknown",
            mime_type="",
            is_valid=False,
            error=validation.error,
        )

    strategy = validation.processing_strategy
    mime_type = validation.mime_type

    try:
        if strategy == "document":
            return _process_document_type(file_name, file_bytes, file_size, mime_type, gemini_client)
        elif strategy == "image":
            return _process_image_type(file_name, file_bytes, file_size, mime_type, gemini_client)
        elif strategy == "docx":
            return _process_docx(file_name, file_bytes, file_size, mime_type)
        elif strategy == "rtf":
            return _process_rtf(file_name, file_bytes, file_size, mime_type)
        elif strategy == "text":
            return _process_text(file_name, file_bytes, file_size, mime_type)
        else:
            return ProcessedDocument(
                file_name=file_name,
                file_size=file_size,
                content_type="unknown",
                mime_type=mime_type,
                is_valid=False,
                error=f"No processing strategy for: {strategy}",
            )
    except Exception as e:
        logger.exception("Error processing %s", file_name)
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="unknown",
            mime_type=mime_type,
            is_valid=False,
            error=f"Processing error: {str(e)}",
        )


def _process_document_type(
    file_name: str,
    file_bytes: bytes,
    file_size: int,
    mime_type: str,
    gemini_client=None,
) -> ProcessedDocument:
    """Process PDF — use native Gemini document understanding.

    For files under INLINE_SIZE_LIMIT, sends inline base64.
    For larger files, uses the Gemini Files API.
    """
    warnings: list[str] = []

    # Large file → use Files API
    if file_size > INLINE_SIZE_LIMIT_BYTES and gemini_client:
        try:
            file_io = io.BytesIO(file_bytes)
            uploaded = gemini_client.files.upload(
                file=file_io,
                config={"mime_type": mime_type},
            )
            return ProcessedDocument(
                file_name=file_name,
                file_size=file_size,
                content_type="document",
                mime_type=mime_type,
                file_uri=uploaded.uri,
                warnings=warnings,
            )
        except Exception as e:
            warnings.append(f"Files API upload failed, falling back to inline: {e}")

    # Inline base64
    data_b64 = base64.b64encode(file_bytes).decode("utf-8")
    return ProcessedDocument(
        file_name=file_name,
        file_size=file_size,
        content_type="document",
        mime_type=mime_type,
        data_base64=data_b64,
        warnings=warnings,
    )


def _process_image_type(
    file_name: str,
    file_bytes: bytes,
    file_size: int,
    mime_type: str,
    gemini_client=None,
) -> ProcessedDocument:
    """Process image files — use Gemini multimodal vision."""
    warnings: list[str] = []

    # Large image → Files API
    if file_size > INLINE_SIZE_LIMIT_BYTES and gemini_client:
        try:
            file_io = io.BytesIO(file_bytes)
            uploaded = gemini_client.files.upload(
                file=file_io,
                config={"mime_type": mime_type},
            )
            return ProcessedDocument(
                file_name=file_name,
                file_size=file_size,
                content_type="image",
                mime_type=mime_type,
                file_uri=uploaded.uri,
                warnings=warnings,
            )
        except Exception as e:
            warnings.append(f"Files API upload failed, falling back to inline: {e}")

    data_b64 = base64.b64encode(file_bytes).decode("utf-8")
    return ProcessedDocument(
        file_name=file_name,
        file_size=file_size,
        content_type="image",
        mime_type=mime_type,
        data_base64=data_b64,
        warnings=warnings,
    )


def _process_docx(
    file_name: str,
    file_bytes: bytes,
    file_size: int,
    mime_type: str,
) -> ProcessedDocument:
    """Extract text from DOCX using python-docx."""
    warnings: list[str] = []

    try:
        from docx import Document
    except ImportError:
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type=mime_type,
            is_valid=False,
            error="python-docx is required for DOCX processing. Install with: pip install python-docx",
        )

    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Preserve heading information
                if para.style and para.style.name and "Heading" in para.style.name:
                    paragraphs.append(f"## {text}")
                else:
                    paragraphs.append(text)

        # Extract table content
        for table in doc.tables:
            table_rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(cells))
            if table_rows:
                paragraphs.append("\n".join(table_rows))

        text_content = "\n\n".join(paragraphs)

        if not text_content.strip():
            warnings.append(
                "No text content could be extracted from this DOCX file. "
                "The document may be image-based or empty."
            )

        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type="text/plain",
            text_content=text_content,
            warnings=warnings,
        )

    except Exception as e:
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type=mime_type,
            is_valid=False,
            error=f"Failed to extract DOCX content: {str(e)}",
        )


def _process_rtf(
    file_name: str,
    file_bytes: bytes,
    file_size: int,
    mime_type: str,
) -> ProcessedDocument:
    """Extract text from RTF using striprtf."""
    warnings: list[str] = []

    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError:
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type=mime_type,
            is_valid=False,
            error="striprtf is required for RTF processing. Install with: pip install striprtf",
        )

    try:
        rtf_content = file_bytes.decode("utf-8", errors="replace")
        text_content = rtf_to_text(rtf_content)

        if not text_content.strip():
            warnings.append("No text content extracted from RTF file.")

        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type="text/plain",
            text_content=text_content,
            warnings=warnings,
        )
    except Exception as e:
        return ProcessedDocument(
            file_name=file_name,
            file_size=file_size,
            content_type="text",
            mime_type=mime_type,
            is_valid=False,
            error=f"Failed to extract RTF content: {str(e)}",
        )


def _process_text(
    file_name: str,
    file_bytes: bytes,
    file_size: int,
    mime_type: str,
) -> ProcessedDocument:
    """Process plain text / markdown files."""
    warnings: list[str] = []

    # Try UTF-8 first, then latin-1 as fallback
    try:
        text_content = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text_content = file_bytes.decode("latin-1")
            warnings.append("File decoded using latin-1 encoding (not UTF-8).")
        except Exception:
            return ProcessedDocument(
                file_name=file_name,
                file_size=file_size,
                content_type="text",
                mime_type=mime_type,
                is_valid=False,
                error="Could not decode text file. Unsupported encoding.",
            )

    if not text_content.strip():
        warnings.append("File appears to be empty or contains only whitespace.")

    return ProcessedDocument(
        file_name=file_name,
        file_size=file_size,
        content_type="text",
        mime_type="text/plain",
        text_content=text_content,
        warnings=warnings,
    )


def build_gemini_input(doc: ProcessedDocument) -> list[dict]:
    """Convert a ProcessedDocument into Gemini API input content list.

    Returns the content list suitable for `client.interactions.create(input=...)`.
    """
    if not doc.is_valid:
        raise ValueError(f"Cannot build input from invalid document: {doc.error}")

    content: list[dict] = []

    if doc.content_type == "document":
        if doc.file_uri:
            content.append({
                "type": "document",
                "uri": doc.file_uri,
                "mime_type": doc.mime_type,
            })
        elif doc.data_base64:
            content.append({
                "type": "document",
                "data": doc.data_base64,
                "mime_type": doc.mime_type,
            })

    elif doc.content_type == "image":
        if doc.file_uri:
            content.append({
                "type": "image",
                "uri": doc.file_uri,
                "mime_type": doc.mime_type,
            })
        elif doc.data_base64:
            content.append({
                "type": "image",
                "data": doc.data_base64,
                "mime_type": doc.mime_type,
            })

    elif doc.content_type == "text":
        content.append({
            "type": "text",
            "text": f"[RESUME DOCUMENT CONTENT — FILE: {doc.file_name}]\n\n{doc.text_content}",
        })

    return content
